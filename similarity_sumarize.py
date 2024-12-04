from flask import Flask, request, render_template, jsonify  
from transformers import BartTokenizer, BartForConditionalGeneration  
import streamlit as st  
from sklearn.feature_extraction.text import TfidfVectorizer  
from sklearn.metrics.pairwise import cosine_similarity  
from wordcloud import WordCloud  
import matplotlib.pyplot as plt  
from docx import Document  
import re  
import nltk  
from nltk.corpus import stopwords  
import pandas as pd  
import time  

# Inisialisasi Flask app  
app = Flask(__name__)  

# Load model dan tokenizer untuk summarization  
tokenizer = BartTokenizer.from_pretrained('facebook/bart-large-cnn')  
model = BartForConditionalGeneration.from_pretrained('facebook/bart-large-cnn')  

# Pastikan untuk mengunduh daftar stop words jika belum dilakukan  
nltk.download('stopwords')  

# Fungsi untuk menghasilkan ringkasan  
def generate_summary(text):  
    inputs = tokenizer.encode("summarize: " + text, return_tensors="pt", max_length=1024, truncation=True)  
    summary_ids = model.generate(inputs, max_length=150, min_length=50, length_penalty=2.0, num_beams=4, early_stopping=True)  
    summary = tokenizer.decode(summary_ids[0], skip_special_tokens=True)  
    return summary  

# Route utama untuk halaman web  
@app.route('/')  
def index():  
    return render_template('index.html')  

# Route untuk proses summarization  
@app.route('/summarize', methods=['POST'])  
def summarize():  
    data = request.form['article']  
    summary = generate_summary(data)  
    return jsonify({"summary": summary})  

# Streamlit untuk kesamaan teks  
def clean_text(text):  
    text = text.lower()  
    text = re.sub(r'[.,\'"]', '', text)  
    text = re.sub(r'\d+', '', text)  
    text = re.sub(r'[^a-zA-Z\s]', '', text)  
    
    stop_words = set(stopwords.words('indonesian'))  
    text = ' '.join(word for word in text.split() if word not in stop_words)  
    
    return text  

def generate_wordcloud(text, title):  
    wordcloud = WordCloud(width=400, height=200, background_color='white').generate(text)  
    plt.figure(figsize=(5, 3))  
    plt.imshow(wordcloud, interpolation='bilinear')  
    plt.axis('off')  
    plt.title(title)  
    st.pyplot(plt)  

# Streamlit application  
st.title("Kalkulator Kesamaan Teks dan Ringkasan")  

input_method = st.radio("Pilih metode input:", ("Input Teks", "Unggah File .docx"))  

text1 = ""  
text2 = ""  

if input_method == "Input Teks":  
    text1 = st.text_area("Masukkan Teks 1:", "", key="text1")  
    text2 = st.text_area("Masukkan Teks 2:", "", key="text2")  

elif input_method == "Unggah File .docx":  
    uploaded_file1 = st.file_uploader("Unggah Dokumen 1 (.docx)", type="docx")  
    uploaded_file2 = st.file_uploader("Unggah Dokumen 2 (.docx)", type="docx")  

    if uploaded_file1 is not None:  
        document1 = Document(uploaded_file1)  
        text1 = "\n".join([para.text for para in document1.paragraphs])  

    if uploaded_file2 is not None:  
        document2 = Document(uploaded_file2)  
        text2 = "\n".join([para.text for para in document2.paragraphs])  

if st.button("Hitung Kesamaan"):  
    if not text1 or not text2:  
        st.warning("Harap masukkan kedua teks.")  
    else:  
        # Ringkas kedua teks  
        summary1 = generate_summary(text1)  
        summary2 = generate_summary(text2)  
        
        st.subheader("Ringkasan Dokumen 1:")  
        st.write(summary1)  
        
        st.subheader("Ringkasan Dokumen 2:")  
        st.write(summary2)  

        # Bersihkan dan hitung kesamaan  
        cleaned_summary1 = clean_text(summary1)  
        cleaned_summary2 = clean_text(summary2)  

        # Hitung kesamaan  
        vectorizer = TfidfVectorizer()  
        tfidf_matrix = vectorizer.fit_transform([cleaned_summary1, cleaned_summary2])  
        cosine_sim = cosine_similarity(tfidf_matrix[0], tfidf_matrix[1])[0][0]  

        st.success(f"Nilai Kesamaan antara ringkasan teks: {cosine_sim:.2f}")  

        # Generate WordCloud untuk ringkasan  
        generate_wordcloud(cleaned_summary1, "WordCloud Ringkasan Dokumen 1")  
        generate_wordcloud(cleaned_summary2, "WordCloud Ringkasan Dokumen 2")  

if __name__ == "__main__":  
    app.run(debug=True)